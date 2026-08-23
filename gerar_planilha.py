# -*- coding: utf-8 -*-
"""
Gerador do Organizador de Declaração de Imposto de Renda.
Tema de identidade: preto, branco e cinza (substituindo o rosa dos exemplos).
Tudo eh construido no Excel: barra de botoes de navegacao, validacoes automaticas,
formulas de resumo e links uteis.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

# ---------- Paleta de cores (preto, branco, cinza) ----------
PRETO   = "1A1A1A"
CINZA_E = "3F3F3F"
CINZA   = "808080"
CINZA_C = "D9D9D9"
BRANCO  = "FFFFFF"
CINZA_F = "F2F2F2"
AZUL    = "1F4E79"

fundo_preto  = PatternFill("solid", fgColor=PRETO)
fundo_cinza_e = PatternFill("solid", fgColor=CINZA_E)
fundo_cinza  = PatternFill("solid", fgColor=CINZA)
fundo_cinza_c = PatternFill("solid", fgColor=CINZA_C)
fundo_cinza_f = PatternFill("solid", fgColor=CINZA_F)
fundo_branco = PatternFill("solid", fgColor=BRANCO)

fonte_branca = Font(name="Calibri", size=11, bold=True, color=BRANCO)
fonte_preta  = Font(name="Calibri", size=11, color=PRETO)
fonte_titulo = Font(name="Calibri", size=18, bold=True, color=BRANCO)
fonte_sub    = Font(name="Calibri", size=10, italic=True, color=BRANCO)
fonte_link   = Font(name="Calibri", size=10, color=AZUL, underline="single")
fonte_peq    = Font(name="Calibri", size=9, italic=True, color=PRETO)

alinhar_centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
alinhar_esq    = Alignment(horizontal="left", vertical="center")
alinhar_esq_w  = Alignment(horizontal="left", vertical="center", wrap_text=True)
alinhar_dir    = Alignment(horizontal="right", vertical="center")

borda_fina = Side(style="thin", color=CINZA)
borda = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

SHEETS = ['Início', 'Dados Pessoais', 'Rendimentos', 'Deduções',
          'Bens e Direitos', 'Resumo', 'Links Úteis']


def barra_navegacao(ws, ncols):
    """Cria uma barra de botoes no topo (linhas 1 e 2) que redireciona entre as abas."""
    n = len(SHEETS)
    base = max(1, ncols // n)
    extra = ncols % n
    col = 1
    for i, nome in enumerate(SHEETS):
        span = base + (1 if i < extra else 0)
        cs, ce = col, col + span - 1
        ws.merge_cells(start_row=1, start_column=cs, end_row=2, end_column=ce)
        cell = ws.cell(row=1, column=cs)
        cell.value = nome
        if nome == ws.title:
            cell.fill = fundo_cinza_e
        else:
            cell.fill = fundo_preto
        cell.font = fonte_branca
        cell.alignment = alinhar_centro
        cell.border = borda
        cell.hyperlink = "#'%s'!A1" % nome
        col = ce + 1
    ws.row_dimensions[1].height = 15
    ws.row_dimensions[2].height = 15


def titulo(ws, texto, ncols, linha=3):
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
    c = ws.cell(row=linha, column=1, value=texto)
    c.fill = fundo_preto
    c.font = fonte_titulo
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = borda
    ws.row_dimensions[linha].height = 28


def cabecalho(ws, linha, valores):
    for i, v in enumerate(valores):
        c = ws.cell(row=linha, column=i + 1, value=v)
        c.fill = fundo_preto
        c.font = fonte_branca
        c.alignment = alinhar_centro
        c.border = borda


def linha_campo(ws, r, label, valor, zebra=False, ph=False):
    a = ws.cell(row=r, column=1, value=label)
    a.fill = fundo_cinza_c if zebra else fundo_branco
    a.font = fonte_preta
    a.alignment = alinhar_esq
    a.border = borda
    b = ws.cell(row=r, column=2, value=valor)
    b.fill = fundo_cinza_c if zebra else fundo_branco
    b.font = Font(name="Calibri", size=11, italic=ph, color=(CINZA if ph else PRETO))
    b.alignment = alinhar_esq
    b.border = borda


wb = openpyxl.Workbook()

# =========================================================
# 1. INICIO
# =========================================================
ini = wb.active
ini.title = "Início"
ini.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [20, 20, 20, 20, 20, 20, 20]):
    ini.column_dimensions[col].width = w
barra_navegacao(ini, 7)
titulo(ini, "ORGANIZADOR DE DECLARAÇÃO DE IMPOSTO DE RENDA", 7, linha=3)
ini.merge_cells("A4:G4")
s = ini["A4"]
s.value = "Ferramenta completa em Excel  •  Tema preto, branco e cinza  •  Use os botões acima para navegar"
s.fill = fundo_cinza_e
s.font = fonte_sub
s.alignment = Alignment(horizontal="center", vertical="center")
s.border = borda

ini.merge_cells("A6:G6")
h = ini["A6"]
h.value = "COMO USAR"
h.fill = fundo_cinza_e
h.font = fonte_branca
h.alignment = alinhar_esq
h.border = borda

passos = [
    "1. Preencha 'Dados Pessoais' com os dados do contribuinte.",
    "2. Lance seus 'Rendimentos' usando a lista de tipos (validação automática).",
    "3. Registre as 'Deduções' e os 'Bens e Direitos' conforme os comprovantes.",
    "4. A aba 'Resumo' consolida tudo automaticamente (fórmulas).",
    "5. Use 'Links Úteis' para acessar sites oficiais durante o preenchimento.",
]
r = 7
for i, p in enumerate(passos):
    ini.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ini.cell(row=r, column=1, value=p)
    c.fill = fundo_cinza_f if (i % 2 == 1) else fundo_branco
    c.font = fonte_preta
    c.alignment = alinhar_esq
    c.border = borda
    r += 1

ini.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=7)
av = ini.cell(row=r + 1, column=1)
av.value = "Dica: os botões no topo de cada aba levam direto à seção desejada."
av.fill = fundo_cinza_c
av.font = fonte_peq
av.alignment = alinhar_centro
av.border = borda

# =========================================================
# 2. DADOS PESSOAIS
# =========================================================
dp = wb.create_sheet("Dados Pessoais")
dp.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [24, 40, 16, 16, 16, 16, 16]):
    dp.column_dimensions[col].width = w
barra_navegacao(dp, 7)
titulo(dp, "DADOS PESSOAIS DO CONTRIBUINTE", 7, linha=3)

dp.merge_cells("A4:G4")
sec = dp["A4"]
sec.value = "Informações cadastrais"
sec.fill = fundo_cinza_e
sec.font = fonte_branca
sec.alignment = alinhar_esq
sec.border = borda

campos = [
    ("Nome completo", "Fulano de Tal Genérico"),
    ("CPF", "000.000.000-00"),
    ("Data de nascimento", "01/01/1980"),
    ("Estado civil", "Casado(a)"),
    ("Endereço", "Rua Exemplo, 123 - Bairro Modelo"),
    ("Cidade / UF", "Cidade Exemplo / EX"),
    ("CEP", "00000-000"),
    ("Telefone", "(00) 00000-0000"),
    ("E-mail", "contribuinte@exemplo.com"),
    ("Tipo de declaração", "Completa"),
    ("Exercício (ano)", "2025"),
    ("Código de agendamento", "00000"),
]
r = 5
for i, (lab, val) in enumerate(campos):
    linha_campo(dp, r, lab, val, zebra=(i % 2 == 1), ph=True)
    r += 1

dp.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=7)
sec2 = dp.cell(row=r + 1, column=1, value="Observações")
sec2.fill = fundo_cinza_e
sec2.font = fonte_branca
sec2.alignment = alinhar_esq
sec2.border = borda
dp.merge_cells(start_row=r + 2, start_column=1, end_row=r + 4, end_column=7)
obs = dp.cell(row=r + 2, column=1)
obs.value = "Use esta área para anotar pendências, documentos recebidos ou lembretes da sua declaração."
obs.fill = fundo_cinza_f
obs.font = fonte_preta
obs.alignment = alinhar_esq_w
obs.border = borda
dp.freeze_panes = "A5"

# =========================================================
# 3. RENDIMENTOS
# =========================================================
rend = wb.create_sheet("Rendimentos")
rend.sheet_view.showGridLines = False
for i, w in enumerate([8, 26, 18, 24, 18, 16, 22]):
    rend.column_dimensions[get_column_letter(i + 1)].width = w
barra_navegacao(rend, 7)
titulo(rend, "CONTROLE DE RENDIMENTOS", 7, linha=3)
headers = ["ID", "Fonte Pagadora", "CNPJ / CPF", "Tipo de Rendimento", "Valor (R$)", "Recebido em", "Observação"]
cabecalho(rend, 4, headers)

tipos = '"Tributável - PJ,Tributável - PF,Isento/Não Tributável,Rendimento de Aplicações,Outros"'
dv = DataValidation(type="list", formula1=tipos, allow_blank=True)
rend.add_data_validation(dv)

N = 40
start = 5
for r in range(start, start + N):
    zebra = ((r - start) % 2 == 1)
    idc = rend.cell(row=r, column=1, value=r - start + 1)
    idc.alignment = alinhar_centro
    idc.fill = fundo_cinza_c if zebra else fundo_branco
    idc.border = borda
    for col in range(2, 8):
        c = rend.cell(row=r, column=col)
        c.border = borda
        c.fill = fundo_cinza_c if zebra else fundo_branco
        c.alignment = alinhar_esq
        if col == 5:
            c.number_format = 'R$ #,##0.00'
            c.alignment = alinhar_dir
    dv.add(rend.cell(row=r, column=4))

tot = start + N
rend.cell(row=tot, column=4, value="TOTAL GERAL").font = fonte_branca
rend.cell(row=tot, column=4).fill = fundo_preto
rend.cell(row=tot, column=4).alignment = alinhar_centro
rend.cell(row=tot, column=4).border = borda
tc = rend.cell(row=tot, column=5, value="=SUM(E%d:E%d)" % (start, tot - 1))
tc.number_format = 'R$ #,##0.00'
tc.font = fonte_branca
tc.fill = fundo_preto
tc.alignment = alinhar_dir
tc.border = borda
for col in [1, 2, 3, 6, 7]:
    rend.cell(row=tot, column=col).fill = fundo_preto
    rend.cell(row=tot, column=col).border = borda

rend.conditional_formatting.add(
    "E%d:E%d" % (start, tot - 1),
    CellIsRule(operator="greaterThan", formula=["10000"],
               fill=PatternFill("solid", fgColor=CINZA),
               font=Font(color=BRANCO, bold=True)))
rend.freeze_panes = "A5"

# =========================================================
# 4. DEDUCOES
# =========================================================
ded = wb.create_sheet("Deduções")
ded.sheet_view.showGridLines = False
for i, w in enumerate([8, 26, 30, 18, 18, 22, 12]):
    ded.column_dimensions[get_column_letter(i + 1)].width = w
barra_navegacao(ded, 7)
titulo(ded, "CONTROLE DE DEDUÇÕES", 7, linha=3)
headers = ["ID", "Categoria", "Descrição", "Valor (R$)", "Documento", "Observação", ""]
cabecalho(ded, 4, headers)

cats = '"Previdência Oficial,Saúde,Mobilidade/Educação,Dependente,Pensão Alimentícia,Doações,Outras"'
dv2 = DataValidation(type="list", formula1=cats, allow_blank=True)
ded.add_data_validation(dv2)

N2 = 30
for r in range(5, 5 + N2):
    zebra = ((r - 5) % 2 == 1)
    idc = ded.cell(row=r, column=1, value=r - 4)
    idc.alignment = alinhar_centro
    idc.fill = fundo_cinza_c if zebra else fundo_branco
    idc.border = borda
    for col in range(2, 8):
        c = ded.cell(row=r, column=col)
        c.border = borda
        c.fill = fundo_cinza_c if zebra else fundo_branco
        c.alignment = alinhar_esq
        if col == 4:
            c.number_format = 'R$ #,##0.00'
            c.alignment = alinhar_dir
    dv2.add(ded.cell(row=r, column=2))

tot2 = 5 + N2
ded.cell(row=tot2, column=3, value="TOTAL DEDUTÍVEL").font = fonte_branca
ded.cell(row=tot2, column=3).fill = fundo_preto
ded.cell(row=tot2, column=3).alignment = alinhar_centro
ded.cell(row=tot2, column=3).border = borda
tc = ded.cell(row=tot2, column=4, value="=SUM(D5:D%d)" % (tot2 - 1))
tc.number_format = 'R$ #,##0.00'
tc.font = fonte_branca
tc.fill = fundo_preto
tc.alignment = alinhar_dir
tc.border = borda
for col in [1, 2, 5, 6, 7]:
    ded.cell(row=tot2, column=col).fill = fundo_preto
    ded.cell(row=tot2, column=col).border = borda
ded.freeze_panes = "A5"

# =========================================================
# 5. BENS E DIREITOS
# =========================================================
bens = wb.create_sheet("Bens e Direitos")
bens.sheet_view.showGridLines = False
for i, w in enumerate([8, 24, 32, 22, 24, 14, 14]):
    bens.column_dimensions[get_column_letter(i + 1)].width = w
barra_navegacao(bens, 7)
titulo(bens, "BENS E DIREITOS (PATRIMÔNIO)", 7, linha=3)
headers = ["ID", "Grupo de Bem", "Descrição", "Situação 31/12 (R$)", "Observação", "", ""]
cabecalho(bens, 4, headers)

grupos = '"Imóveis,Veículos,Aplicações Financeiras,Participação em Empresas,Bens Móveis,Outros"'
dv3 = DataValidation(type="list", formula1=grupos, allow_blank=True)
bens.add_data_validation(dv3)

N3 = 25
for r in range(5, 5 + N3):
    zebra = ((r - 5) % 2 == 1)
    idc = bens.cell(row=r, column=1, value=r - 4)
    idc.alignment = alinhar_centro
    idc.fill = fundo_cinza_c if zebra else fundo_branco
    idc.border = borda
    for col in range(2, 8):
        c = bens.cell(row=r, column=col)
        c.border = borda
        c.fill = fundo_cinza_c if zebra else fundo_branco
        c.alignment = alinhar_esq
        if col == 4:
            c.number_format = 'R$ #,##0.00'
            c.alignment = alinhar_dir
    dv3.add(bens.cell(row=r, column=2))

tot3 = 5 + N3
bens.cell(row=tot3, column=3, value="TOTAL DE BENS").font = fonte_branca
bens.cell(row=tot3, column=3).fill = fundo_preto
bens.cell(row=tot3, column=3).alignment = alinhar_centro
bens.cell(row=tot3, column=3).border = borda
tc = bens.cell(row=tot3, column=4, value="=SUM(D5:D%d)" % (tot3 - 1))
tc.number_format = 'R$ #,##0.00'
tc.font = fonte_branca
tc.fill = fundo_preto
tc.alignment = alinhar_dir
tc.border = borda
for col in [1, 2, 5, 6, 7]:
    bens.cell(row=tot3, column=col).fill = fundo_preto
    bens.cell(row=tot3, column=col).border = borda
bens.freeze_panes = "A5"

# =========================================================
# 6. RESUMO
# =========================================================
res = wb.create_sheet("Resumo")
res.sheet_view.showGridLines = False
for col, w in zip("ABC", [3, 44, 22]):
    res.column_dimensions[col].width = w
barra_navegacao(res, 3)
titulo(res, "RESUMO CONSOLIDADO", 3, linha=3)

linhas = [
    ("Total de Rendimentos", "=Rendimentos!E%d" % (start + N), 'R$ #,##0.00'),
    ("Total de Deduções", "=Deduções!D%d" % (5 + N2), 'R$ #,##0.00'),
    ("Total de Bens", "=Bens e Direitos!D%d" % (5 + N3), 'R$ #,##0.00'),
]
r = 5
for i, (lab, formula, fmt) in enumerate(linhas):
    zebra = (i % 2 == 1)
    a = res.cell(row=r, column=2, value=lab)
    a.fill = fundo_cinza_c if zebra else fundo_branco
    a.font = fonte_preta
    a.alignment = alinhar_esq
    a.border = borda
    c = res.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    c.fill = fundo_cinza_c if zebra else fundo_branco
    c.font = fonte_preta
    c.alignment = alinhar_dir
    c.border = borda
    r += 1

res.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
sec = res.cell(row=r, column=2, value="Cálculo simplificado (exemplo genérico)")
sec.fill = fundo_cinza_e
sec.font = fonte_branca
sec.alignment = alinhar_esq
sec.border = borda
r += 1

def res_linha(r, label, valor, fmt=None, bold=False):
    a = res.cell(row=r, column=2, value=label)
    a.fill = fundo_cinza_c
    a.font = Font(bold=bold, color=PRETO)
    a.alignment = alinhar_esq
    a.border = borda
    c = res.cell(row=r, column=3, value=valor)
    if fmt:
        c.number_format = fmt
    c.fill = fundo_cinza_c
    c.font = Font(bold=bold, color=PRETO)
    c.alignment = alinhar_dir
    c.border = borda

res_linha(r, "Base de cálculo (Rendimentos - Deduções)", "=C5-C6", 'R$ #,##0.00'); r += 1
res_linha(r, "Alíquota estimada (%)", 15, '0"%"'); r += 1
a = res.cell(row=r, column=2, value="Imposto estimado (referência)")
a.fill = fundo_cinza_c
a.font = Font(bold=True, color=PRETO)
a.alignment = alinhar_esq
a.border = borda
c = res.cell(row=r, column=3, value="=C9*C10/100")
c.number_format = 'R$ #,##0.00'
c.fill = fundo_cinza_c
c.font = Font(bold=True, color=PRETO)
c.alignment = alinhar_dir
c.border = borda
r += 2

res.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
av = res.cell(row=r, column=2)
av.value = "Atenção: cálculo genérico e educativo. Consulte a legislação vigente e um contador."
av.fill = fundo_cinza_c
av.font = fonte_peq
av.alignment = alinhar_centro
av.border = borda
res.freeze_panes = "A5"

# =========================================================
# 7. LINKS UTEIS
# =========================================================
lk = wb.create_sheet("Links Úteis")
lk.sheet_view.showGridLines = False
for col, w in zip("ABC", [3, 42, 78]):
    lk.column_dimensions[col].width = w
barra_navegacao(lk, 3)
titulo(lk, "LINKS ÚTEIS (ACESSO RÁPIDO)", 3, linha=3)

lk.merge_cells("B4:C4")
sec = lk["B4"]
sec.value = "Sites Oficiais"
sec.fill = fundo_cinza_e
sec.font = fonte_branca
sec.alignment = alinhar_esq
sec.border = borda

oficiais = [
    ("Receita Federal - IRPF", "https://www.gov.br/receitafederal/pt-br/assuntos/meu-imposto-de-renda"),
    ("Programa IRPF (download)", "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/download/pgd"),
    ("Consulta à Restituição", "https://www.gov.br/receitafederal/pt-br/servicos/consultar-restituicao"),
    ("e-CAC (centro virtual)", "https://www.gov.br/receitafederal/pt-br/centrais-de-atendimento/cac"),
    ("Tabelas Práticas da IE", "https://www.gov.br/receitafederal/pt-br/assuntos/meu-imposto-de-renda/irpf/2025/tabelas"),
]
r = 5
for i, (nome, url) in enumerate(oficiais):
    zebra = (i % 2 == 1)
    a = lk.cell(row=r, column=2, value=nome)
    a.fill = fundo_cinza_c if zebra else fundo_branco
    a.font = fonte_preta
    a.alignment = alinhar_esq
    a.border = borda
    c = lk.cell(row=r, column=3, value=url)
    c.hyperlink = url
    c.font = fonte_link
    c.fill = fundo_cinza_f if zebra else fundo_branco
    c.alignment = alinhar_esq
    c.border = borda
    r += 1

r += 1
lk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
sec = lk.cell(row=r, column=2, value="Materiais e Apoio (Desafio DIO)")
sec.fill = fundo_cinza_e
sec.font = fonte_branca
sec.alignment = alinhar_esq
sec.border = borda
r += 1
apoio = [
    ("GitHub Quick Start (DIO)", "https://github.com/digitalinnovationone/github-quick-start"),
    ("Documentação do GitHub", "https://docs.github.com/pt"),
    ("GitHub Markdown Guide", "https://docs.github.com/pt/get-started/writing-on-github/getting-started-with-writing-and-formatting-on-github"),
    ("Azure - VM Windows", "https://learn.microsoft.com/pt-br/azure/virtual-machines/windows/quick-create-portal"),
    ("Guia IRPF - Gov.br", "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/perguntas-e-respostas/irpf"),
]
for i, (nome, url) in enumerate(apoio):
    zebra = (i % 2 == 1)
    a = lk.cell(row=r, column=2, value=nome)
    a.fill = fundo_cinza_c if zebra else fundo_branco
    a.font = fonte_preta
    a.alignment = alinhar_esq
    a.border = borda
    c = lk.cell(row=r, column=3, value=url)
    c.hyperlink = url
    c.font = fonte_link
    c.fill = fundo_cinza_f if zebra else fundo_branco
    c.alignment = alinhar_esq
    c.border = borda
    r += 1
lk.freeze_panes = "A5"

wb.save("Organizador_IRPF.xlsx")
print("Arquivo Organizador_IRPF.xlsx gerado com sucesso.")
print("Abas:", wb.sheetnames)
